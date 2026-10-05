"""Build the Review 2 and Review 3 decks on the Amrita template.

Both start from docs/SafeScape_Review_Presentation.pptx so the template's master (crimson
footer, logo) and its title/team slide carry over; every other slide is removed and
rebuilt here in the same style. Numbers come from the same places as the report
(19_build_report.py): reports/metrics/*_v2.json and the served model bundles.

  Review 2 — one model per member, demonstrated with the dataset
  Review 3 — every model integrated with the UI, results compared, UI tested

Usage (from the repo root): .venv/Scripts/python.exe scripts/20_build_review_decks.py
"""
import copy
import importlib
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
report = importlib.import_module("19_build_report")
MODELS, LABELS, bundle_kb, pct = report.MODELS, report.LABELS, report.bundle_kb, report.pct

BASE = ROOT / "docs" / "SafeScape_Review_Presentation.pptx"
FIG = ROOT / "reports" / "figures"

RED = RGBColor(0xFF, 0x00, 0x00)
HEAD = RGBColor(0xC0, 0x00, 0x00)
ACCENT = RGBColor(0xB6, 0x11, 0x4D)
INK = RGBColor(0x1A, 0x1A, 0x2E)
MUTED = RGBColor(0x5A, 0x5A, 0x6B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ZEBRA = RGBColor(0xF6, 0xF1, 0xF1)

OWNER = {"mfcc_cnn": "Amruth Rohan KR", "logmel_crnn": "Harish Venkat VS", "transformer": "Harish Venkat VS"}
ARCH_FIG = {"mfcc_cnn": "mfcc_cnn_architecture.png", "logmel_crnn": "crnn_architecture.png",
            "transformer": "transformer_architecture.png"}


# ----------------------------------------------------------------------------- deck plumbing
def fresh_deck(review_label):
    """The template deck with only its title/team slide left, tagged with the review."""
    prs = Presentation(str(BASE))
    ids = prs.slides._sldIdLst
    for sld in list(ids)[1:]:
        prs.part.drop_rel(sld.rId)
        ids.remove(sld)
    title = prs.slides[0]
    tb = title.shapes.add_textbox(Inches(9.9), Inches(0.42), Inches(3.0), Inches(0.5))
    p = tb.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    r = p.add_run()
    r.text = review_label
    r.font.size = Pt(18)
    r.font.bold = True
    r.font.name = "Cambria"
    r.font.color.rgb = HEAD
    return prs


def new_slide(prs, title, tag):
    s = prs.slides.add_slide(prs.slides[0].slide_layout)
    tb = s.shapes.add_textbox(Inches(0.5), Inches(0.28), Inches(10.0), Inches(0.58))
    r = tb.text_frame.paragraphs[0].add_run()
    r.text = title
    r.font.size = Pt(26 if len(title) > 44 else 30)
    r.font.bold = True
    r.font.name = "Cambria"
    r.font.color.rgb = RED
    tb = s.shapes.add_textbox(Inches(10.6), Inches(0.34), Inches(2.23), Inches(0.5))
    p = tb.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    r = p.add_run()
    r.text = tag
    r.font.size = Pt(10.5)
    r.font.color.rgb = MUTED
    return s


def text(slide, x, y, w, h, blocks):
    """blocks: list of (kind, text) with kind in heading / body / bullet / verdict."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    for kind, t in blocks:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        r = p.add_run()
        r.text = ("▪  " + t) if kind == "bullet" else t
        r.font.name = "Calibri"
        if kind == "heading":
            r.font.size, r.font.bold, r.font.color.rgb = Pt(11), True, HEAD
            p.space_before = Pt(6)
        elif kind == "verdict":
            r.font.size, r.font.bold, r.font.color.rgb = Pt(12), True, ACCENT
            p.space_before = Pt(6)
        else:
            r.font.size, r.font.color.rgb = Pt(12.5 if kind == "body" else 12), INK
            p.space_after = Pt(4)
    return tb


def table(slide, headers, rows, x, y, w, col_w=None, fs=11, row_h=0.32):
    shape = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(x), Inches(y), Inches(w),
                                   Inches(row_h * (len(rows) + 1)))
    t = shape.table
    if col_w:
        for i, cw in enumerate(col_w):
            t.columns[i].width = Inches(cw)
    for r_i, row in enumerate([headers] + rows):
        for c_i, v in enumerate(row):
            cell = t.cell(r_i, c_i)
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.text = ""
            run = cell.text_frame.paragraphs[0].add_run()
            run.text = str(v)
            run.font.size = Pt(fs)
            run.font.name = "Calibri"
            cell.fill.solid()
            if r_i == 0:
                run.font.bold = True
                run.font.color.rgb = WHITE
                cell.fill.fore_color.rgb = HEAD
            else:
                run.font.color.rgb = INK
                cell.fill.fore_color.rgb = ZEBRA if r_i % 2 == 0 else WHITE
    return t


def picture(slide, img, x, y, max_w, max_h):
    from PIL import Image
    w, h = Image.open(img).size
    scale = min(max_w / w, max_h / h)
    pw, ph = w * scale, h * scale
    return slide.shapes.add_picture(str(img), Inches(x + (max_w - pw) / 2), Inches(y), Inches(pw), Inches(ph))


def metric_row(k, M):
    m = M[k]
    return [MODELS[k], OWNER[k], pct(m["accuracy"]), f"{m['macro_f1']:.3f}", pct(m["distress_recall"]),
            f"{m['params']:,}", f"{bundle_kb(k):.0f} KB", f"{m['latency_ms']:.2f} ms"]


METRIC_HEAD = ["model", "owner", "accuracy", "macro-F1", "distress recall", "params", "served size", "CPU / window"]


# ----------------------------------------------------------------------------- shared slides
def s_dataset(prs, tag):
    src, split = report.dataset_tables()
    s = new_slide(prs, "Dataset — sources and class mapping", tag)
    table(s, ["source"] + LABELS + ["total"],
          [[i] + [int(src.loc[i, l]) for l in LABELS] + [int(src.loc[i].sum())] for i in src.index]
          + [["total"] + [int(src[l].sum()) for l in LABELS] + [int(src.values.sum())]],
          0.6, 1.05, 12.1, col_w=[2.5] + [1.6] * 5 + [1.6], fs=11)
    text(s, 0.6, 4.05, 6.0, 2.6, [
        ("heading", "WHERE EACH CLASS COMES FROM"),
        ("bullet", "ESC-50 — glass, horn, siren, clock alarm, ambient scenes"),
        ("bullet", "Kaggle scream set + RAVDESS angry/fearful speech → distress_call"),
        ("bullet", "UrbanSound8K — car_horn → horn_skid, siren → alarm, 7 urban classes as hard-negative ambience"),
        ("bullet", "DataSEC + Freesound CC0 → extra glass_break"),
    ])
    text(s, 6.9, 4.05, 5.8, 2.6, [
        ("heading", "DESIGN DECISIONS"),
        ("bullet", "UrbanSound8K gun_shot excluded: a real hazard with no class here — labelling it "
                   "ambience would teach the model gunfire is normal"),
        ("bullet", "16 kHz mono, denoised, 1 s windows with 50% overlap, near-silent windows dropped"),
        ("verdict", f"2,919 recordings → {int(split.values.sum()):,} windows"),
    ])


def s_split(prs, tag):
    _, split = report.dataset_tables()
    s = new_slide(prs, "Dataset — leakage-free split and duplicate audit", tag)
    table(s, ["class", "train", "val", "test (leak-free)", "removed from test"],
          [[l, int(split.loc[l, "train"]), int(split.loc[l, "val"]), int(split.loc[l, "test"]),
            int(split.loc[l].get("test_dup", 0))] for l in LABELS]
          + [["total", int(split["train"].sum()), int(split["val"].sum()), int(split["test"].sum()),
              int(split["test_dup"].sum())]],
          0.6, 1.05, 6.6, col_w=[1.8, 1.1, 1.1, 1.4, 1.2], fs=11.5)
    text(s, 7.6, 1.0, 5.1, 5.4, [
        ("heading", "SPLIT"),
        ("body", "70 / 15 / 15, stratified, at the source-file level — every window of a recording "
                 "stays on one side."),
        ("heading", "AUDIT — FOUND DURING THIS REVIEW"),
        ("body", "Hashing every source file showed the Kaggle scream set ships the same recordings in two "
                 "folders. 230 test windows (all distress_call) were byte-identical to training audio."),
        ("body", "Fix without retraining: flag them (18_flag_duplicate_sources.py) and report every number "
                 "on the remaining 2,990 windows."),
        ("verdict", "Impact was small (CRNN 89.8% → 89.7%), but every headline number is now on audio no "
                    "model has heard."),
    ])


def s_features(prs, tag):
    s = new_slide(prs, "Feature pipeline — shared by all three models", tag)
    table(s, ["step", "setting", "why"],
          [["resample + mono", "16 kHz", "speech/hazard energy is below 8 kHz; smaller input"],
           ["denoise", "non-stationary spectral gating", "suppress steady background before features"],
           ["window", "1.0 s, 50% hop", "short enough for near-real-time, long enough for a scream"],
           ["energy gate", "RMS ≥ 1e-4", "never train or predict on silence"],
           ["STFT", "512-pt FFT, 10 ms hop → 101 frames", "10 ms resolution keeps glass transients"],
           ["MFCC (CNN)", "40 coefficients", "compact cepstral summary of the spectrum"],
           ["log-mel (CRNN, Transformer)", "64 mel bands, dB", "keeps spectral detail the DCT discards"],
           ["normalise", "per window, zero mean / unit var", "loudness-invariant input"],
           ["augment (train only)", "SpecAugment masks + feature noise", "regularise the rare classes"]],
          0.6, 1.05, 12.1, col_w=[2.6, 3.6, 5.9], fs=11.5)


def s_model(prs, k, M, tag):
    m = M[k]
    s = new_slide(prs, f"{MODELS[k]} — {OWNER[k]}", tag)
    picture(s, FIG / ARCH_FIG[k], 0.6, 0.95, 12.1, 2.55)
    hp = {"mfcc_cnn": [("learning rate", "1.25e-3"), ("batch / dropout", "64 / 0.25"),
                       ("selection", "Optuna, val macro recall"), ("epochs", "25, early stop at 17")],
          "logmel_crnn": [("learning rate", "8.56e-4"), ("batch / dropout", "16 / 0.27"),
                          ("GRU", "128 units × 2 directions"), ("compression", "dynamic int8 → 516 KB")],
          "transformer": [("teacher", "v2 CRNN (frozen)"), ("α / T", "0.5 / 4"),
                          ("learning rate / batch", "1e-3 / 32"), ("epochs", "25, early stop at 20")]}[k]
    table(s, ["hyper-parameter", "value"], [list(r) for r in hp], 0.6, 3.75, 4.6, col_w=[2.0, 2.6], fs=11)
    table(s, ["class", "precision", "recall", "F1"],
          [[l, f"{m['per_class'][l]['precision']:.2f}", f"{m['per_class'][l]['recall']:.2f}",
            f"{m['per_class'][l]['f1-score']:.2f}"] for l in LABELS],
          5.5, 3.75, 4.4, col_w=[1.6, 0.95, 0.95, 0.9], fs=11)
    text(s, 10.15, 3.7, 2.6, 2.8, [
        ("heading", "LEAK-FREE TEST"),
        ("body", f"accuracy {pct(m['accuracy'])}"),
        ("body", f"macro-F1 {m['macro_f1']:.3f}"),
        ("body", f"distress recall {pct(m['distress_recall'])}"),
        ("body", f"{m['params']:,} params"),
        ("body", f"{bundle_kb(k):.0f} KB · {m['latency_ms']:.2f} ms"),
    ])


def s_cms(prs, M, tag, title):
    s = new_slide(prs, title, tag)
    for i, k in enumerate(MODELS):
        picture(s, ROOT / M[k]["confusion_png"], 0.5 + i * 4.15, 1.0, 4.0, 3.5)
        text(s, 0.5 + i * 4.15, 4.55, 4.0, 0.4, [("heading", f"{MODELS[k]} · {OWNER[k]}")])
    text(s, 0.6, 5.1, 12.1, 1.4, [
        ("body", "Rows are the true class. The dominant error for every model is ambience predicted as a "
                 "hazard — the price of class-weighted training and recall-floored calibration, which "
                 "deliberately prefer a false alarm to a missed scream."),
    ])


def s_curves(prs, tag):
    s = new_slide(prs, "Training — validation macro recall per epoch", tag)
    picture(s, FIG / "training_curves_v2.png", 0.6, 1.0, 7.8, 4.6)
    text(s, 8.7, 1.1, 4.0, 5.0, [
        ("heading", "HOW EACH MODEL WAS SELECTED"),
        ("bullet", "Checkpoint = best validation macro recall, not lowest loss — recall on rare hazards is what matters"),
        ("bullet", "Early stopping after 6 epochs without improvement"),
        ("bullet", "Class-weighted cross-entropy (1 / class frequency) for all three"),
        ("bullet", "Then a per-class logit bias fitted on validation, with every hazard's recall floored"),
    ])


def s_compare(prs, M, tag):
    s = new_slide(prs, "Three-model comparison — leak-free test split", tag)
    table(s, METRIC_HEAD, [metric_row(k, M) for k in MODELS], 0.6, 1.05, 12.1,
          col_w=[2.2, 2.0, 1.2, 1.2, 1.5, 1.3, 1.4, 1.3], fs=12, row_h=0.4)
    picture(s, FIG / "architecture_comparison.png", 0.6, 2.85, 7.2, 3.7)
    c, t, m = M["logmel_crnn"], M["transformer"], M["mfcc_cnn"]
    text(s, 8.1, 2.85, 4.6, 3.8, [
        ("heading", "READING THE TABLE"),
        ("bullet", f"CRNN leads every quality metric: macro-F1 {c['macro_f1']:.3f} vs "
                   f"{t['macro_f1']:.3f} / {m['macro_f1']:.3f}"),
        ("bullet", f"Transformer: {t['params'] / c['params']:.0%} of the CRNN's parameters, "
                   f"{c['latency_ms'] / t['latency_ms']:.0f}× faster per window"),
        ("bullet", "MFCC-CNN: smallest and fastest, weakest on glass_break"),
        ("verdict", "All three are far inside the 1 s real-time budget, so hazard recall decides: "
                    "the quantized CRNN ships as the default."),
    ])


# ----------------------------------------------------------------------------- Review 2
def build_review2(M):
    prs = fresh_deck("Review 2")
    tag = "Review 2 · 5 marks"
    s = new_slide(prs, "Review 2 — one model per member, with the dataset", tag)
    table(s, ["member", "model", "features", "what is demonstrated"],
          [["Amruth Rohan KR", "MFCC-CNN", "40 MFCCs × 101 frames",
            "train/eval on the v2 dataset, per-class results, live classification of held-out clips"],
           ["Harish Venkat VS", "log-mel CRNN", "64 log-mel × 101 frames",
            "same, plus calibration and int8 quantization of the champion"],
           ["Harish Venkat VS", "Distilled Transformer", "64 log-mel × 101 frames",
            "distillation from the CRNN teacher, same evaluation"]],
          0.6, 1.1, 12.1, col_w=[2.2, 2.2, 2.4, 5.3], fs=12, row_h=0.5)
    text(s, 0.6, 3.4, 12.1, 3.0, [
        ("heading", "COMMON GROUND"),
        ("bullet", "One dataset (v2), one source-file-disjoint split, one leak-free test set of 2,990 windows"),
        ("bullet", "Same metrics for every model: accuracy, macro-F1, distress-call recall, size, CPU latency"),
        ("bullet", "Each model is calibrated the same way, so they are compared at equivalent operating points"),
    ])
    s_dataset(prs, tag)
    s_split(prs, tag)
    s_features(prs, tag)
    for k in MODELS:
        s_model(prs, k, M, tag)
    s_curves(prs, tag)
    s_cms(prs, M, tag, "Results — confusion matrices per member")
    s_compare(prs, M, tag)
    s = new_slide(prs, "Live demonstration", tag)
    text(s, 0.6, 1.05, 12.1, 5.4, [
        ("heading", "EACH MEMBER RUNS THEIR OWN MODEL ON THE HELD-OUT TEST SET"),
        ("body", "cd scripts;  $env:SAFESCAPE_PROC_DIR = \"data/processed_v2\""),
        ("body", "python 10_evaluate.py --arch mfcc_cnn --ckpt ../models/checkpoints/mfcc_cnn_v2.pt "
                 "--calibration ../models/exported/mfcc_cnn/calibration.json "
                 "--manifest ../data/processed_v2/windows_manifest_dedup.csv --cpu-only        (Amruth)"),
        ("body", "python 10_evaluate.py --arch logmel_crnn --hidden-size 128 --ckpt ../models/checkpoints/"
                 "logmel_crnn_v2.pt --calibration ../models/exported/logmel_crnn/calibration.json "
                 "--manifest ../data/processed_v2/windows_manifest_dedup.csv --cpu-only        (Harish)"),
        ("body", "python 10_evaluate.py --arch transformer --ckpt ../models/checkpoints/transformer_v2.pt "
                 "--calibration ../models/exported/transformer/calibration.json "
                 "--manifest ../data/processed_v2/windows_manifest_dedup.csv --cpu-only        (Harish)"),
        ("heading", "THEN THE SAME MODEL ON A SINGLE CLIP"),
        ("body", "python run_demo.py  →  http://127.0.0.1:8124/compare  →  pick a held-out test clip; "
                 "each member's card shows their model's verdict, probabilities and latency."),
    ])
    out = ROOT / "docs" / "SafeScape_Review2.pptx"
    prs.core_properties.title = "SafeScape — Review 2"
    prs.save(out)
    return out


# ----------------------------------------------------------------------------- Review 3
def build_review3(M, grid):
    prs = fresh_deck("Review 3")
    tag = "Review 3 · 6 marks"

    s = new_slide(prs, "Integration — every model behind one interface", tag)
    picture(s, FIG / "system_architecture.png", 0.6, 0.95, 12.1, 2.7)
    table(s, ["endpoint", "used by", "does"],
          [["POST /predict?model=<key>", "mobile app", "one model (default CRNN) — unchanged from Review 1"],
           ["POST /predict/all", "dashboard", "one clip → all three models, with per-model latency"],
           ["GET /models", "dashboard", "test metrics, owner and served size per model"],
           ["GET /testclips", "dashboard", "20 held-out clips with ground truth"]],
          0.6, 3.85, 12.1, col_w=[3.0, 1.8, 7.3], fs=11.5)
    text(s, 0.6, 5.65, 12.1, 0.8, [
        ("verdict", "Each model is a self-contained bundle (weights, label map, pre-processing config, "
                    "calibration); all share the same multi-window serving rule, so differences are architectural."),
    ])

    s = new_slide(prs, "Comparison dashboard — same clip, every model", tag)
    picture(s, FIG / "ui_compare_live.png", 0.6, 0.95, 8.0, 5.6)
    text(s, 8.9, 1.0, 3.9, 5.5, [
        ("heading", "INPUTS"),
        ("bullet", "record 2 s from the microphone"),
        ("bullet", "upload any audio file"),
        ("bullet", "pick one of 20 held-out test clips"),
        ("heading", "PER MODEL"),
        ("bullet", "label, confidence, full probability bars"),
        ("bullet", "server inference time"),
        ("bullet", "✓ / ✗ against ground truth when known"),
        ("verdict", "No external libraries or CDN — runs offline on the demo laptop."),
    ])

    s = new_slide(prs, "UI testing — 20 held-out clips through the interface", tag)
    picture(s, FIG / "ui_compare_uitest.png", 0.5, 0.95, 6.6, 5.6)
    correct = {k: sum(g[2][k] == g[1] for g in grid) for k in MODELS}
    per_cls = [[l] + [f"{sum(g[2][k] == l for g in grid if g[1] == l)}/4" for k in MODELS] for l in LABELS]
    table(s, ["class", "MFCC-CNN", "CRNN", "Transformer"],
          per_cls + [["total"] + [f"{correct[k]}/{len(grid)}" for k in MODELS]],
          7.4, 1.05, 5.3, col_w=[1.7, 1.2, 1.2, 1.2], fs=11.5)
    text(s, 7.4, 3.5, 5.3, 3.0, [
        ("heading", "WHAT THE TEST SHOWS"),
        ("bullet", "Hazards: 47 of 48 model×clip calls correct"),
        ("bullet", "Ambience is the weak spot: the serving rule lets any hazard window decide the clip — "
                   "a deliberate false-alarm bias"),
        ("bullet", "Clips are cut around their loudest second, which makes ambience harder than real use"),
        ("verdict", "Next: require two consecutive hazard windows before alerting."),
    ])

    s = new_slide(prs, "Functional and edge-case tests", tag)
    table(s, ["test", "expected", "result"],
          [["Hazard in first or second half of a clip", "detected", "pass — all 4 hazard classes"],
           ["Ambience followed by 1 s of silence", "ambience", "fail — distress_call 0.77 (same with the "
            "earlier model; pre-processing)"],
           ["2 s of silence", "ambience, every model", "pass — energy gate"],
           ["0.33 s clip", "every model answers", "pass — zero-padded window"],
           ["Unknown model key", "HTTP 404, server stays up", "pass"],
           ["No microphone access (LAN http)", "record disabled with reason", "pass — upload/test clips still work"],
           ["Existing mobile app on the new server", "unchanged", "pass — 80/99 smoke test, all 5 classes"],
           ["First request latency", "inference only", "pass — models warmed at startup (1.3 s → 47 ms)"]],
          0.6, 1.05, 12.1, col_w=[4.0, 3.0, 5.1], fs=11.5, row_h=0.42)
    text(s, 0.6, 5.0, 12.1, 1.4, [
        ("body", "Found and fixed by UI testing: the server originally judged only the last second of each "
                 "2 s clip, so a glass break in the first half was missed. Every window is now classified and "
                 "the regression test guards it."),
    ])

    s = new_slide(prs, "Offline metrics, as served by the dashboard", tag)
    picture(s, FIG / "ui_compare_offline.png", 0.6, 0.95, 12.1, 5.7)

    s_compare(prs, M, tag)

    s = new_slide(prs, "Per-class comparison", tag)
    table(s, ["class"] + [f"{report.SHORT[k]} {x}" for k in MODELS for x in ("P", "R", "F1")],
          [[l] + [f"{M[k]['per_class'][l][x]:.2f}" for k in MODELS for x in ("precision", "recall", "f1-score")]
           for l in LABELS],
          0.6, 1.05, 12.1, fs=11.5, row_h=0.4)
    text(s, 0.6, 3.75, 12.1, 2.6, [
        ("heading", "DISCUSSION"),
        ("bullet", "glass_break separates the models most (F1 " + " / ".join(
            f"{M[k]['per_class']['glass_break']['f1-score']:.2f}" for k in MODELS) +
            "): a transient of a few hundred ms that the CRNN's time axis catches and pooling blurs"),
        ("bullet", "alarm and horn_skid are strong for every model since UrbanSound8K added ~5× more data"),
        ("bullet", "distress_call recall stays above 80% for all three — the class the proposal prioritises"),
    ])

    s_cms(prs, M, tag, "Confusion matrices — leak-free test split")

    s = new_slide(prs, "Limitations and next steps", tag)
    text(s, 0.6, 1.05, 12.1, 5.4, [
        ("bullet", "Distress audio is acted / emotional speech — a consented, self-recorded India-context set is "
                   "the most valuable next step"),
        ("bullet", "169 validation windows still duplicate training audio — re-split on content hashes"),
        ("bullet", "Ambience false alarms at clip level — require agreement between consecutive windows"),
        ("bullet", "Ambience followed by silence reads as distress_call — likely per-window normalisation, "
                   "not yet confirmed"),
        ("bullet", "Inference runs on a laptop over LAN — export to TorchScript/ONNX for on-phone inference"),
        ("bullet", "Phone microphone needs HTTPS — certificate or native wrapper"),
        ("verdict", "Deliverables: working app + dashboard, three trained models, case study report, "
                    "code at github.com/Hash-486/SafeScape"),
    ])
    out = ROOT / "docs" / "SafeScape_Review3.pptx"
    prs.core_properties.title = "SafeScape — Review 3"
    prs.save(out)
    return out


def main():
    M = report.load_metrics()
    report.training_curves()
    grid = report.ui_test_grid()
    for p in (build_review2(M), build_review3(M, grid)):
        print(f"wrote {p}")


if __name__ == "__main__":
    main()
